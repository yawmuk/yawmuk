#!/usr/bin/env python3
"""Build the «يومك» submission deck from the official challenge template.

Output: docs/deck/yawmuk-deck.template.pptx  (keeps {{TOKENS}}; fill.mjs replaces them)

Needs python-pptx (pip install python-pptx). Run from anywhere:
    python3 docs/deck/build.py [path/to/project_template.pptx]

Every {{TOKEN}} is written as a single run so fill.mjs can replace it with a plain string
replace. Every screenshot is a named picture ("SLOT:<NAME>") so fill.mjs can swap its bytes.
"""
import copy
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..', '..', '..'))  # /Users/.../islamicaich
TEMPLATE = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'project_template.pptx')
OUT = os.path.join(HERE, 'yawmuk-deck.template.pptx')
IMG = os.path.join(HERE, 'img')

NAVY = RGBColor(0x12, 0x18, 0x3F)
CARD = RGBColor(0x1C, 0x23, 0x5C)
PURPLE = RGBColor(0x61, 0x50, 0xEA)
TURQ = RGBColor(0x2E, 0xF2, 0xC2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xC9, 0xCF, 0xF2)
FONT = 'Readex Pro'

prs = Presentation(TEMPLATE)
src = list(prs.slides)  # src[n-1] == template slide n

# ---------------------------------------------------------------- helpers

def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError(sid)


def _set_para(p, text):
    runs = p.runs
    if not runs:
        r = p.add_run()
        r.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    # drop line breaks left inside the paragraph
    for br in p._p.findall(qn('a:br')):
        p._p.remove(br)


def set_text(slide, sid, lines):
    """Replace a shape's text with `lines` (str or list), cloning the first paragraph's style."""
    if isinstance(lines, str):
        lines = [lines]
    tf = shape(slide, sid).text_frame
    paras = tf.paragraphs
    proto = copy.deepcopy(paras[0]._p)
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    _set_para(tf.paragraphs[0], lines[0])
    last = tf.paragraphs[0]._p
    for line in lines[1:]:
        np_ = copy.deepcopy(proto)
        last.addnext(np_)
        last = np_
    for p, line in zip(tf.paragraphs, lines):
        _set_para(p, line)


def set_paras(slide, sid, texts):
    """Replace paragraph-by-paragraph, keeping each paragraph's own style."""
    tf = shape(slide, sid).text_frame
    paras = tf.paragraphs
    texts = list(texts) + [None] * (len(paras) - len(texts))
    assert len(paras) == len(texts), (sid, len(paras), [p.text for p in paras])
    for p, t in zip(paras, texts):
        if t is None:
            continue
        if t == '' and not p.runs:
            continue
        _set_para(p, t)


def delete(slide, *sids):
    for sid in sids:
        el = shape(slide, sid)._element
        el.getparent().remove(el)


def rtl_para(p, align=PP_ALIGN.RIGHT):
    p.alignment = align
    p._p.get_or_add_pPr().set('rtl', '1')


def textbox(slide, x, y, w, h, paras, name=None, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.RIGHT):
    """paras: list of (text, size_pt, color, bold)."""
    tb = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, Pt(2))
    for i, (text, size, color, bold) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        rtl_para(p, align)
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color
        f.name = FONT
        rpr = r._r.get_or_add_rPr()
        cs = etree.SubElement(rpr, qn('a:cs'))
        cs.set('typeface', FONT)
    return tb


def card(slide, x, y, w, h, fill=CARD, line=PURPLE, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    s.adjustments[0] = 0.08
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def arrow_left(slide, x, y, w=26, h=22):
    a = slide.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Pt(x), Pt(y), Pt(w), Pt(h))
    a.fill.solid()
    a.fill.fore_color.rgb = TURQ
    a.line.fill.background()
    a.shadow.inherit = False
    return a


SLOTS = {}


def picture(slide, ph_sid, path, slot):
    ph = shape(slide, ph_sid)
    aspect = round(ph.width / ph.height, 4)
    geo = (ph.left, ph.top, ph.width, ph.height)
    pic = ph.insert_picture(path)
    pic.left, pic.top, pic.width, pic.height = geo
    pic.name = f'SLOT:{slot}'
    SLOTS[slot] = {
        'aspect': aspect,
        'default': os.path.relpath(path, HERE),
    }
    return pic


NOTES = []


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
    title = next((sh.text_frame.text for sh in slide.placeholders
                  if sh.placeholder_format.type in (1, 3) and sh.has_text_frame), '')
    NOTES.append((title.replace('\n', ' '), text))


SECTION = {}

# ================================================================ 1 · Cover (template 8)
s = src[7]
set_text(s, 486, '«يومك»')
set_text(s, 487, ['لعبة ثلاثية الأبعاد: يوم كامل في بلدة واحدة،',
                  'بأدلة موثّقة ومرشد ذكي مقيّد بالمصادر'])
set_text(s, 489, 'فريق Lemonada · المسار 03: التجارب التفاعلية')
set_text(s, 490, '6 أكتوبر 2026')
# widen the team line to the left so it does not wrap
sh = shape(s, 489)
sh.left, sh.width = Pt(760), Pt(566)
notes(s, 'السلام عليكم. نحن فريق Lemonada، ونقدّم «يومك»: لعبة ثلاثية الأبعاد في المتصفح، يعيش فيها المتعلّم يوماً عادياً في بلدة واحدة، ويرى كيف يتعامل المسلمون مع مواقفه اليومية، ثم يقرأ الدليل الموثّق. (15 ثانية)')

# ================================================================ 2 · Problem (template 11)
s = src[10]
set_text(s, 537, 'المشكلة')
set_text(s, 538, 'أسئلة يومية حقيقية، وإجابات لا تناسب المتعلّم')
set_text(s, 539, 'قرض عقاري بفائدة، حفلة عمل فيها خمر، محفظة مفقودة، سحب خيري: كيف يتصرف المسلم؟ ولماذا؟')
set_text(s, 541, 'فتاوى نصية طويلة')
set_text(s, 542, 'تشرح الحكم، لكنها لا تجعل المتعلّم يعيش الموقف ولا تبيّن التصرف العملي.')
set_text(s, 543, 'مقاطع قصيرة بلا مصادر')
set_text(s, 544, 'سريعة الانتشار، لكنها نادراً ما تذكر الآية أو الحديث ودرجته.')
set_text(s, 545, 'روبوتات محادثة عامة')
set_text(s, 546, 'تجيب عن كل شيء، وقد تنسب نصاً لا مصدر له، وتختزل الخلاف، ولا تحيل إلى عالم.')
set_text(s, 547, 'يوم المسلم مجزّأ')
set_text(s, 548, 'المواقيت والمصحف والأذكار والمعاملات في تطبيقات متفرقة لا تربطها رحلة واحدة.')
notes(s, 'المشكلة: من يتعرّف على الإسلام، وكثير من المسلمين في الغرب، تواجههم أسئلة يومية: قرض عقاري، حفلة عمل فيها خمر، محفظة مفقودة. المتاح إما فتوى نصية طويلة، أو مقطع قصير بلا مصدر، أو روبوت محادثة عام يجيب عن كل شيء وقد ينسب نصاً لا مصدر له ولا يحيل إلى عالم. ويوم المسلم نفسه موزّع على تطبيقات متفرقة. (25 ثانية)')

# ================================================================ 3 · Solution (template 13)
s = src[12]
set_text(s, 573, 'الحل')
set_text(s, 574, 'بلدة واحدة يعيش فيها المتعلّم يوماً كاملاً')
set_text(s, 577, 'ما الذي نراه؟')
set_text(s, 578, 'بلدة ثلاثية الأبعاد متصلة: البيت والعمل والكلية والشارع والحفلات والمسجد والمصرف الإسلامي، يتنقّل بينها اللاعب مشياً.')
set_text(s, 579, 'ما الذي يهم المستخدم؟')
set_text(s, 580, '18 موقفاً ببطاقات حكم موثّقة، ومواقيت صلاة حيّة مع الأذان، ومصحف مرتّل مع معاني الآيات، وأذكار موثّقة، وحاسبة زكاة.')
picture(s, 576, os.path.join(IMG, 'solution.png'), 'SCREENSHOT_TOWN')
notes(s, 'الحل: «يومك» بلدة واحدة متصلة. يخرج اللاعب من البيت إلى العمل والكلية والشارع، ويدخل المسجد والمصرف الإسلامي. في كل محطة موقف حقيقي، وبعده بطاقة حكم بالدليل. وفي المسجد مواقيت صلاة حيّة وأذان، ومصحف مرتّل مع معاني الآيات، وأذكار الصباح والمساء من مصادر موثّقة. وفي المصرف مستشار للمعاملات وحاسبة زكاة. (25 ثانية)')

# ================================================================ 4 · How it works (template 24)
s = src[23]
set_text(s, 818, 'آلية العمل')
set_text(s, 819, 'رحلة المتعلّم: من الاهتمام إلى الفهم')
set_text(s, 820, 'لا حساب، ولا سؤال عن المعتقد، ولا تخزين لما يكتبه اللاعب')
set_paras(s, 822, ['01', 'خطة الرحلة', 'يكتب اللاعب أو يقول ما يهمه', 'فيرتّب المخطط الذكي مواقفه'])
set_paras(s, 823, ['02', 'عِش الموقف', 'يتجوّل ويحاور الشخصيات', 'ويختار كيف يتصرف آدم'])
set_paras(s, 824, ['03', 'بطاقة الحكم', 'آيات وأحاديث موثّقة حرفياً', 'والمذاهب الأربعة والمرشد'])
set_paras(s, 825, ['04', 'تحقّق وتابع', 'سؤال فهم ثم المحطة التالية', 'وقياس الفهم قبل وبعد'])
notes(s, 'آلية العمل في أربع خطوات: يقول اللاعب ما يهمه، كتابةً أو بصوته، فيرتّب المخطط الذكي رحلته من مكتبة المواقف المراجعة. يعيش الموقف ويختار كيف يتصرف. تظهر بطاقة الحكم: الآيات والأحاديث حرفياً بأرقامها ودرجاتها، والمذاهب الأربعة جنباً إلى جنب. ثم سؤال فهم، والقياس قبل اللعب وبعده. ولا نسأل عن المعتقد ولا نخزّن ما يكتبه اللاعب. (25 ثانية)')

# ================================================================ 5 · AI in detail (template 27, custom diagram)
s = src[26]
set_text(s, 860, 'الذكاء الاصطناعي')
set_text(s, 861, 'مرشد يجيب من المصادر المراجعة فقط، أو يمتنع ويحيل')
set_text(s, 862, 'توليد مقيّد بالاسترجاع (Gemini 3 Flash على Vertex AI عبر الخادم) · المفتاح لا يصل إلى المتصفح · بديل حتمي عند أي عطل')
delete(s, 864, 865, 866, 867, 868, 869, 870, 871, 872, 873, 874, 875)

steps = [
    ('01', 'سؤال بالصوت أو الكتابة', 'تعرّف على الكلام ونطق للإجابة عبر Web Speech'),
    ('02', 'فرز قبل النموذج', 'سؤال عن حالة شخصية = إحالة مباشرة إلى عالم'),
    ('03', 'استرجاع مقيّد', 'من المقاطع المراجعة وبطاقة الحكم الحالية فقط'),
    ('04', 'Gemini بمخرجات JSON', 'معرّفات المقاطع قائمة مغلقة؛ لا نص ديني حر'),
    ('05', 'مدقّق آلي', 'إجابة موسومة بمصادرها، أو امتناع وإحالة'),
]
x0, y0, gap, bw, bh = 1326, 300, 30, 223, 178
for i, (num, title, body) in enumerate(steps):
    x = x0 - (i + 1) * bw - i * gap
    card(s, x, y0, bw, bh, name=f'AI step {num}')
    textbox(s, x + 12, y0 + 10, bw - 24, bh - 20,
            [(num, 22, TURQ, True), (title, 19, WHITE, True), (body, 16, SOFT, False)],
            name=f'AI step {num} text')
    if i < len(steps) - 1:
        arrow_left(s, x - gap + 2, y0 + bh / 2 - 11)

cards = [
    ('مخطط الرحلة', 'يرتّب المواقف الثمانية عشر بحسب اهتمام اللاعب؛ معرّفات معروفة فقط، وعند الفشل ترتيب حتمي من الكتالوج نفسه.'),
    ('حلقة المراجعة البشرية', 'كل سؤال يمتنع عنه المرشد أو يحيله يصل إلى لوحة الخبراء الشرعيين؛ وما يعتمدونه يصبح مقطعاً يجيب منه المرشد.'),
    ('التقييم الآلي', 'الحالات الاختبارية الاثنتا عشرة في الحزمة العلمية تُشغَّل آلياً على المرشد، ويُقاس الامتناع والإحالة والإسناد.'),
]
cw, ch, cy = 385, 175, 515
for i, (title, body) in enumerate(cards):
    x = 1326 - (i + 1) * cw - i * 30
    card(s, x, cy, cw, ch, fill=RGBColor(0x16, 0x1C, 0x4A), line=None, name=f'AI card {i+1}')
    textbox(s, x + 18, cy + 14, cw - 36, ch - 28,
            [(title, 20, TURQ, True), (body, 17, WHITE, False)], name=f'AI card {i+1} text')
notes(s, 'الذكاء الاصطناعي هنا ليس روبوتاً عاماً. المرشد يستقبل السؤال بالصوت أو الكتابة. قبل أي نموذج نفرز الأسئلة الشخصية ونحيلها إلى عالم. ثم نسترجع المقاطع المراجعة فقط، ونرسلها إلى Gemini على Google Cloud من الخادم، ويُلزَم بإرجاع JSON يذكر معرّفات المقاطع التي اعتمد عليها. مدقّق آلي يرفض أي معرّف لم نرسله، وأي نص يشبه آية أو حديثاً أو رابطاً، فإن فشل امتنع المرشد وأحال. وكل سؤال يمتنع عنه يصل إلى لوحة خبراء شرعيين، وما يعتمدونه يصبح مصدراً جديداً. وعند غياب المفتاح أو أي عطل تعمل اللعبة ببديل حتمي. (40 ثانية)')

# ================================================================ 6 · Reliability (template 17)
s = src[16]
set_text(s, 639, 'الموثوقية والسلامة العلمية')
set_text(s, 640, 'كل محتوى مصنّف بمستويات الحزمة العلمية الأربعة')
set_text(s, 641, 'تحقّق بنفسك: npm test · npm run audit · docs/SOURCES.md · docs/QA_BANK.md')
tbl = [sh for sh in s.shapes if sh.has_table][0].table
# visual order is right→left = XML columns 3,2,1,0
rows = [
    ('المستوى', 'المعالجة في «يومك»', 'مثال', 'الحارس'),
    ('أ · نص أصلي مستقر', 'آيات بنص مجمع الملك فهد، وأحاديث برقمها ودرجتها؛ تُعرض حرفياً ولا يولّدها النموذج', 'بطاقة الحكم، المصحف، الأذكار', 'مطابقة آلية حرفاً بحرف'),
    ('ب · شرح واستدلال', 'من المادة المراجعة مع المرجع؛ وشرحنا موسوم «توضيحي، ليس نصاً شرعياً»', '«بكلمات بسيطة»، إجابة المرشد', 'معرّفات مصادر إلزامية'),
    ('ج · مسائل خلافية', 'المذاهب الأربعة جنباً إلى جنب، مع نطاق الحكم «قول الجمهور» وبيان الخلاف', 'الرهن العقاري، المصافحة', 'اختبار المستويات'),
    ('د · فتوى أو حالة شخصية', 'لا حكم مستقل: معلومة عامة وإحالة إلى عالم، والسؤال يصل إلى لوحة الخبراء', '«أنا في بلد كذا، هل يجوز لي…؟»', 'فرز قبل النموذج + مراجعة بشرية'),
]
for r, vals in enumerate(rows):
    for c, v in zip((3, 2, 1, 0), vals):
        cell = tbl.cell(r, c)
        _set_para(cell.text_frame.paragraphs[0], v)
set_text(s, 644, 'المصادر: نص مجمع الملك فهد وترجماته عبر quranenc.com، والحديث ودرجته من dorar.net، والشاملة؛ وغيرها موسوم للمراجعة')
notes(s, 'الموثوقية: صنّفنا كل محتوى بمستويات الحزمة العلمية الأربعة. المستوى أ نصوص تُعرض حرفياً ولا يولّدها النموذج أبداً، ونطابقها آلياً مع مصدرها. المستوى ب شرح من المادة المراجعة، وشرحنا موسوم بأنه ليس نصاً شرعياً. المستوى ج مسائل خلافية نعرض فيها المذاهب الأربعة ونطاق الحكم. والمستوى د حالات شخصية لا نحكم فيها، بل نحيل، ويصل السؤال إلى لوحة الخبراء. ويمكن لأي محكّم إعادة الفحص بالأوامر المكتوبة أعلى الشريحة. الأحكام مسودات أعدّها الذكاء الاصطناعي وتنتظر مراجعة عالم، ولا ندّعي غير ذلك. (35 ثانية)')

# ================================================================ 7 · Added value (template 28)
s = src[27]
set_text(s, 883, 'القيمة المضافة')
set_text(s, 884, '«يومك» مقابل روبوت محادثة عام')
set_text(s, 885, 'أساس المقارنة: الحالات الاختبارية الاثنتا عشرة في الحزمة العلمية · النتيجة المقيسة: {{EVAL_COMPARISON}}')
set_text(s, 887, '«يومك»')
set_text(s, 888, 'تجربة الموقف ثم الدليل، ومرشد مقيّد بالمصادر المراجعة')
set_text(s, 889, 'روبوت محادثة عام')
set_text(s, 890, 'إجابة نصية سريعة عن أي سؤال، من معرفة النموذج العامة')
set_paras(s, 891, ['يستشهد بمعرّفات مقاطع تحقّقنا منها، أو يمتنع', '', 'يحيل الحالة الشخصية إلى عالم، والسؤال يصل إلى خبراء', '', 'رحلة مخصّصة داخل بلدة، بلا حساب ولا تتبّع للمعتقد'])
set_paras(s, 892, ['قد ينسب نصاً أو حديثاً دون مصدر يمكن تتبّعه', '', 'قد يجيب عن الحالة الشخصية بحكم مستقل', '', 'لا تجربة ولا تسلسل تعليمي ولا قياس للفهم'])
notes(s, 'القيمة المضافة مقارنة ببديل محدد: روبوت محادثة عام. الروبوت العام يجيب عن أي سؤال من معرفته، وقد ينسب نصاً بلا مصدر، وقد يجيب عن الحالة الشخصية بحكم مستقل. «يومك» يجعل المتعلّم يعيش الموقف أولاً، ثم يجيبه المرشد من مصادر تحقّقنا منها فقط أو يمتنع ويحيل. وقِسنا الفرق آلياً على الحالات الاثنتي عشرة في الحزمة العلمية: {{EVAL_COMPARISON}}. (30 ثانية)')

# ================================================================ 8 · Screenshots (template 26)
s = src[25]
set_text(s, 843, 'لقطات من المشروع')
set_text(s, 844, 'من داخل اللعبة ولوحة الخبراء')
picture(s, 846, os.path.join(IMG, 'gallery1.png'), 'SCREENSHOT_QURAN')
picture(s, 851, os.path.join(IMG, 'gallery2.png'), 'SCREENSHOT_DASHBOARD')
picture(s, 852, os.path.join(IMG, 'gallery3.png'), 'SCREENSHOT_PRAYER')
set_text(s, 847, 'المصحف المرتّل مع معاني الآيات')
set_text(s, 848, 'لوحة الخبراء: فرز آلي ثم إجابة عالم')
set_text(s, 849, 'مواقيت الصلاة الحيّة والأذان')
notes(s, 'لقطات من المشروع: المصحف المرتّل في المسجد، يستمع اللاعب للتلاوة آيةً آية ويقرأ معناها، ولوحة المراجعة الشرعية حيث يصل سؤال اللاعب ويُفرز آلياً، ثم يجيب عنه خبير ويعتمده. وهنا لوحة مواقيت الصلاة في المسجد مع الأذان. وكلها في الرابط الحي: {{LIVE_URL}}. (20 ثانية)')

# ================================================================ 9 · Results (template 15)
s = src[14]
set_text(s, 604, 'النتائج')
set_text(s, 605, 'ما قِسناه، بأرقام يمكن إعادة توليدها')
set_text(s, 606, 'الدراسة القبلية/البعدية مع الفئة المستهدفة: {{STUDY_RESULTS}}')
set_text(s, 608, '{{TESTS_PASS}}')
set_text(s, 609, 'اختباراً آلياً ناجحاً')
set_text(s, 610, ['من أصل {{TESTS_TOTAL}}', 'npm test'])
set_text(s, 611, '{{EVAL_PASS}}')
set_text(s, 612, 'نجاح فحوص الموقع الحي')
set_text(s, 613, ['682 من 695 · الحالات الـ12 × 3', '0 إخفاق غير آمن'])
set_text(s, 614, '{{SOURCES_VERIFIED}}')
set_text(s, 615, 'مصدراً مطابَقاً آلياً')
set_text(s, 616, ['من أصل {{SOURCES_TOTAL}}', 'docs/SOURCES.md'])
set_text(s, 617, '{{STUDY_N}}')
set_text(s, 618, 'مشاركاً في الدراسة')
set_text(s, 619, ['متوسط تحسّن الفهم', '{{STUDY_GAIN}}'])
set_text(s, 620, 'المصدر: المستودع وصفحة النتائج الحيّة results.html، بتاريخ {{RESULTS_DATE}}')
notes(s, 'النتائج، وكلها قابلة لإعادة التوليد: {{TESTS_PASS}} اختباراً آلياً ناجحاً. على الحالات الاثنتي عشرة في الحزمة: {{EVAL_PASS}}. {{SOURCES_VERIFIED}} مصدراً طوبق آلياً حرفاً بحرف. ودراسة قبلية وبعدية مع الفئة المستهدفة: {{STUDY_RESULTS}}. النتائج الحيّة في صفحة results، ولا نعرض رقماً لم نقِسه. (30 ثانية)')

# ================================================================ 10 · Operational realism (template 20)
s = src[19]
set_text(s, 710, 'واقعية التشغيل')
set_text(s, 711, 'تكلفة منخفضة، وبديل لكل اعتماد')
set_text(s, 712, 'خادم Node واحد على Cloud Run يقدّم اللعبة ودوال الذكاء الاصطناعي')
set_text(s, 714, 'التكلفة لكل 1000 لاعب')
set_text(s, 715, 'تقدير أعلى ≈ 45 دولاراً لكل ألف لاعب (محسوب بأسعار Claude؛ التشغيل الحالي Gemini Flash على Vertex AI)، والأزرار والإجابات المعدّة مجانية. القياس الفعلي: {{COST_MEASURED}}')
set_text(s, 716, 'الاعتماديات والبدائل')
set_text(s, 717, 'عند تعذّر النموذج (Gemini 3 Flash ثم 2.5 Flash): مخطط حتمي وإجابات معدّة مسبقاً. المواقيت تُحسب محلياً بمكتبة adhan-js. المصحف من api.quran.com والتلاوة من everyayah.com.')
set_text(s, 718, 'صيانة المحتوى')
set_text(s, 719, 'المحتوى في ملفات JSON منفصلة عن الكود؛ والخبراء يجيبون ويراجعون الأحكام من لوحة التحكم؛ والفحص الآلي يعيد مطابقة النصوص.')
set_text(s, 720, 'خطة التبنّي')
set_text(s, 721, 'مراكز إسلامية ودعوية تستضيف اللعبة في الأيام المفتوحة ودروس المسلمين الجدد؛ الكود مفتوح المصدر.')
notes(s, 'واقعية التشغيل: خادم واحد على Cloud Run. التكلفة التقديرية نحو 45 دولاراً لكل ألف لاعب، والأجزاء المعدّة مسبقاً مجانية. لكل اعتماد بديل: إن غاب النموذج يعمل مخطط حتمي وإجابات معدّة، والمواقيت تُحسب محلياً. المحتوى منفصل عن الكود، ويصونه الخبراء من لوحة التحكم. والتبنّي عبر المراكز الإسلامية في الأيام المفتوحة ودروس المسلمين الجدد. (30 ثانية)')

# ================================================================ 11 · Built vs next (template 18)
s = src[17]
set_text(s, 652, 'ما أُنجز وما يأتي')
set_text(s, 653, 'خطة الاستمرار')
set_text(s, 654, 'المراحل الثلاث الأولى منجزة وتعمل في الرابط الحي (4–6 أكتوبر)، والأخيرتان مقترح لاحق')
stages = [
    (656, 657, 658, 'أُنجز: المواقف', '4–5 أكتوبر', '18 موقفاً في 6 أماكن، وبطاقات حكم موثّقة'),
    (660, 661, 662, 'أُنجز: البلدة', '6 أكتوبر', 'بلدة موحّدة، مسجد، أذان، مصحف، أذكار، مصرف'),
    (664, 665, 666, 'أُنجز: الضبط', '6 أكتوبر', 'مرشد صوتي، لوحة الخبراء، تقييم آلي، دراسة'),
    (668, 669, 670, 'التالي: المراجعة', 'الشهر الأول', 'عالم مؤهل يراجع كل حكم ويوقّعه'),
    (672, 673, 674, 'التالي: التوسع', 'الأشهر 2–6', 'شريك تبنٍّ، مواقف جديدة، لغات أخرى'),
]
for a, b, c, t1, t2, t3 in stages:
    set_text(s, a, t1)
    set_text(s, b, t2)
    set_text(s, c, t3)
notes(s, 'ما أُنجز وما يأتي، بفصل واضح: المراحل الثلاث الأولى منجزة وتعمل في الرابط الحي: المواقف الثمانية عشر، والبلدة الموحّدة بالمسجد والمصرف، والمرشد الصوتي ولوحة الخبراء والتقييم. التالي: مراجعة عالم مؤهل لكل حكم وتوقيعه في الشهر الأول، ثم شريك تبنٍّ ومواقف ولغات جديدة. (20 ثانية)')

# ================================================================ 12 · Thanks (template 31)
s = src[30]
set_text(s, 915, '«يومك» · فريق Lemonada: أبوبكر أبوشام، محمد المبارك')
set_text(s, 916, ['{{LIVE_URL}}', 'github.com/B4r4k4/yawmuk  |  {{VIDEO_URL}}'])
for p in shape(s, 916).text_frame.paragraphs:
    p._p.get_or_add_pPr().set('rtl', '0')
notes(s, 'شكراً لكم. جرّبوا «يومك» في الرابط الحي، والمستودع مفتوح بكل اختباراته ومصادره. نسعد بأسئلتكم. (5 ثوانٍ)')

# ---------------------------------------------------------------- order & drop
order = [8, 11, 13, 24, 27, 17, 28, 26, 15, 20, 18, 31]
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)
keep = [ids[n - 1] for n in order]
for el in ids:
    sldIdLst.remove(el)
    if el not in keep:
        prs.part.drop_rel(el.get(qn('r:id')))
for el in keep:
    sldIdLst.append(el)

prs.save(OUT)
with open(os.path.join(HERE, 'slots.json'), 'w', encoding='utf-8') as f:
    json.dump(SLOTS, f, ensure_ascii=False, indent=2)
# notes are recorded in build order; re-sort into presentation order
order_titles = {n: i for i, n in enumerate(order)}
build_order = [8, 11, 13, 24, 27, 17, 28, 26, 15, 20, 18, 31]
with open(os.path.join(HERE, 'SPEAKER_NOTES.md'), 'w', encoding='utf-8') as f:
    f.write('<div dir="rtl">\n\n# «يومك» — نص العرض (5 دقائق)\n\n')
    f.write('> مولَّد من `build.py` (ملاحظات المتحدث داخل الشرائح). التوقيت التقريبي بين قوسين؛ المجموع ≈ 5 دقائق.\n')
    f.write('> الرموز `{{…}}` تُستبدل عند التعبئة بـ `fill.mjs`؛ اقرأ القيمة الفعلية من الشريحة.\n\n')
    for i, (n, (title, text)) in enumerate(sorted(zip(build_order, NOTES), key=lambda t: order_titles[t[0]]), 1):
        f.write(f'## {i}. {title}\n\n{text}\n\n')
    f.write('</div>\n')
print('wrote', OUT, len(prs.slides._sldIdLst), 'slides')
